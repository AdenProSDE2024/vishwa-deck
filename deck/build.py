#!/usr/bin/env python3
"""Build the deck from per-page files.

    python3 build.py            # deck.html -> PDF -> pptx (+ presenter set)
    python3 build.py --no-notes # skip the internal presenter/notes outputs
    python3 build.py --presenter # also inline the ?presenter=1 overlay (internal only)

Pages live one to a file:
    pages/NN-name.html      presented slides, numbered 01..N in filename order
    appendix/aN-name.html   appendix slides, numbered A1..AM in filename order

Numbering is assigned HERE, not written in the files. Each file carries {{NUM}} in its
footer span (and, for appendix files, in its "Appendix {{NUM}}" eyebrow). Reorder or insert
a page by renaming files — page numbers, appendix labels and their eyebrows stay in sync.
That is the whole point: hand-numbering broke three times before this existed.

Cross-references to appendix pages are written as {{REF:a2-business-model}} and resolved to
the label that file ends up with, so moving an appendix page cannot leave a dangling "A5".
"""
import glob, os, re, subprocess, sys, shutil, time

DATE   = '2026-09-09'
CHROME = '/Applications/Google Chrome.app/Contents/MacOS/Google Chrome'
MIRROR = os.path.expanduser('~/Documents/Dev/whoIs/vishwa-pitch')
WPM, SLIDE_CHANGE = 100.0, 2.0
SYSFONTS = ('Lucida','Menlo','Helvetica','Arial','Times','Courier')
BANNED   = re.compile(r'\[NEED|CONFIDENTIAL|TODO|TBD|DRAFT\b|placeholder', re.I)
HERE = os.path.dirname(os.path.abspath(__file__)); os.chdir(HERE)


VC = '--vc' in sys.argv          # longer investor version: 5 extra body slides
# pages-vc/ holds the slides that only the VC cut carries. These two already exist as
# appendix pages in the short deck, so the VC cut promotes them into the body instead of
# duplicating them — they are removed from the appendix in that build.
VC_PROMOTE = ['appendix/a8-competition.html', 'appendix/a2-business-model.html']
# Pages the VC cut swaps out for its own version. The four-minute deck keeps the amount
# off the ask page (Plug and Play is the wrong room for it); the VC cut puts it back.
# Files in pages-vc/ whose name starts with _ are replacements, not extra slides.
VC_REPLACE = {'pages/11-ask.html': 'pages-vc/_11-ask-vc.html'}


def collect():
    pages = sorted(glob.glob('pages/*.html'))
    apx   = sorted(glob.glob('appendix/*.html'), key=lambda p: int(re.search(r'/a(\d+)', p).group(1)))
    if VC:
        extra = [f for f in glob.glob('pages-vc/*.html')
                 if not os.path.basename(f).startswith('_')]
        pages = sorted(pages + extra, key=lambda p: os.path.basename(p))
        # promoted pages go straight after the extra VC slides. Work this out from `extra`
        # only: a VC_REPLACE file also lives in pages-vc/, and counting it here pushed the
        # promoted pages to the very end of the deck.
        at = max(pages.index(f) for f in extra) + 1
        pages[at:at] = VC_PROMOTE
        for orig, repl in VC_REPLACE.items():
            if orig in pages and os.path.exists(repl):
                pages[pages.index(orig)] = repl
        apx = [f for f in apx if f not in VC_PROMOTE]
    labels = {}
    for i, f in enumerate(pages, start=1): labels[f] = f'{i:02d}'
    for i, f in enumerate(apx,   start=1): labels[f] = f'A{i}'
    ref = {os.path.splitext(os.path.basename(f))[0]: labels[f] for f in apx}
    # a promoted page is no longer in the appendix, so "appendix {{REF:x}}" pointing at it
    # would dangle. Map it to its body slide number and let assemble() drop the word.
    promoted = {os.path.splitext(os.path.basename(f))[0]: labels[f]
                for f in pages if f.startswith('appendix/')}
    return pages, apx, labels, ref, promoted


def assemble():
    pages, apx, labels, ref, promoted = collect()
    out = [open('head.html', encoding='utf8').read()]
    for f in pages + apx:
        s = open(f, encoding='utf8').read()
        if VC and f in VC_PROMOTE:
            # promoted into the body: drop the "Appendix {{NUM}} · " eyebrow prefix,
            # otherwise it reads "Appendix 13" in the middle of the deck
            s = re.sub(r'(<div class="eyebrow">)Appendix \{\{NUM\}\}\s*(?:&nbsp;)?\s*(?:·|&#183;)\s*(?:&nbsp;)?\s*',
                       lambda m: m.group(1) + '', s, count=1)
        s = s.replace('{{NUM}}', labels[f])
        s = re.sub(r'(?:appendix|Appendix)\s+\{\{REF:([a-z0-9\-]+)\}\}',
                   lambda m: f'slide {promoted[m.group(1)]}' if m.group(1) in promoted
                             else f'appendix {ref.get(m.group(1), "A?")}', s)
        s = re.sub(r'\{\{REF:([a-z0-9\-]+)\}\}',
                   lambda m: promoted.get(m.group(1)) or ref.get(m.group(1), 'A?'), s)
        out.append(s)
    if '--presenter' in sys.argv and os.path.exists('presenter-inline.html'):
        out.append('\n' + open('presenter-inline.html', encoding='utf8').read())
    out.append('\n</body></html>')
    html = ''.join(out)
    open('deck.html', 'w', encoding='utf8').write(html)
    return html, pages, apx, ref


def render(html):
    # Chrome 152's headless mode writes the PDF in a few seconds but then does NOT
    # exit (it keeps GCM/updater work alive), so subprocess.run() waits forever.
    # Launch it, wait for the file to appear and stop growing, then kill it.
    # Render to a per-run temp path: two Chromes sharing one --print-to-pdf target
    # corrupt it, and an orphan from a killed build will do that silently.
    stem = f'Vishwa_VC_Pitch_{DATE}' if VC else f'Vishwa_Pitch_{DATE}'
    pdf, tmp = f'{stem}.pdf', f'/tmp/vp_{os.getpid()}.pdf'
    prof = f'/tmp/cr_build_{os.getpid()}'
    shutil.rmtree(prof, ignore_errors=True)
    if os.path.exists(tmp): os.remove(tmp)

    p = subprocess.Popen([CHROME, '--headless', '--disable-gpu', f'--user-data-dir={prof}',
                          '--no-first-run', '--no-default-browser-check', '--no-pdf-header-footer',
                          f'--print-to-pdf={tmp}', f'file://{HERE}/deck.html'],
                         stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    size, stable, deadline, done = -1, 0, time.time() + 300, False
    try:
        while time.time() < deadline:
            if p.poll() is not None:
                done = True; break                  # older Chrome: exits on its own
            cur = os.path.getsize(tmp) if os.path.exists(tmp) else -1
            stable = stable + 1 if cur == size and cur > 0 else 0
            size = cur
            if stable >= 3:
                done = True; break                  # unchanged for ~3s: finished writing
            time.sleep(1)
    finally:
        p.terminate()
        try: p.wait(timeout=10)
        except subprocess.TimeoutExpired: p.kill()

    if not done:
        sys.exit('Chrome never finished writing the PDF (300s).')
    if not os.path.exists(tmp) or os.path.getsize(tmp) < 50_000:
        sys.exit(f'Chrome produced no usable PDF at {tmp}')
    os.replace(tmp, pdf)
    shutil.rmtree(prof, ignore_errors=True)
    return pdf

def main():
    html, pages, apx, ref = assemble()
    pdf = render(html)
    import pymupdf
    d = pymupdf.open(pdf)

    bad = {sp['font'] for i in range(d.page_count)
           for b in d[i].get_text('dict')['blocks'] for l in b.get('lines', [])
           for sp in l['spans'] if any(x in sp['font'] for x in SYSFONTS)}
    leaks = [(i+1, ln.strip()[:60]) for i in range(d.page_count)
             for ln in d[i].get_text().split('\n') if BANNED.search(ln)]
    dangling = sorted(set(re.findall(r'\{\{REF:[a-z0-9\-]+\}\}|appendix A\?', html)))
    balance = None if html.count('<div') == html.count('</div>') else \
              f"<div>={html.count('<div')} </div>={html.count('</div>')}"

    os.makedirs('slides_png', exist_ok=True)
    for f in glob.glob('slides_png/*.png'): os.remove(f)
    pngs = []
    for i in range(d.page_count):
        p = f'slides_png/slide_{i+1:02d}.png'
        d[i].get_pixmap(matrix=pymupdf.Matrix(2.0, 2.0)).save(p); pngs.append(p)

    from pptx import Presentation
    from pptx.util import Inches
    want_notes = '--no-notes' not in sys.argv
    notes = {}
    if want_notes:
        try:
            sys.path.insert(0, HERE)
            import notes as _n
            byfile = getattr(_n, 'NOTES_BY_FILE', {})
            def unwrap(s):
                # notes.py is hard-wrapped for reading in the file. A notes pane turns every
                # one of those line breaks into its own paragraph, which chops sentences in
                # half. Join wrapped lines; keep blank lines and bracketed headers as breaks.
                out = []
                for para in s.strip().split('\n\n'):
                    lines = [l.strip() for l in para.split('\n') if l.strip()]
                    buf = []
                    for l in lines:
                        if l.startswith(('[', '->')):
                            if buf: out.append(' '.join(buf)); buf = []
                            out.append(l)
                        else:
                            buf.append(l)
                    if buf: out.append(' '.join(buf))
                    out.append('')
                return '\n'.join(out).strip()

            notes = {i: unwrap(byfile[os.path.splitext(os.path.basename(f))[0]])
                     for i, f in enumerate(pages, start=1)
                     if os.path.splitext(os.path.basename(f))[0] in byfile}
        except Exception as e:
            print('note: notes.py unavailable —', e)

    made = []
    stem = f'Vishwa_VC_Pitch_{DATE}' if VC else f'Vishwa_Pitch_{DATE}'
    for with_notes, name in ((False, f'{stem}.pptx'),
                             (True,  f'{stem}_presenter.pptx')):
        if with_notes and not notes: continue
        prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        for i, p in enumerate(pngs, start=1):
            sl = prs.slides.add_slide(prs.slide_layouts[6])
            sl.shapes.add_picture(p, 0, 0, width=prs.slide_width, height=prs.slide_height)
            if with_notes and notes.get(i):
                sl.notes_slide.notes_text_frame.text = notes[i].strip()
        prs.save(name)
        # python-pptx writes the notesMaster relationship but not the matching
        # <p:notesMasterIdLst>, and PowerPoint/Keynote refuse the file over it.
        import fix_pptx_notes
        fix_pptx_notes.repair(name)
        import os as _os
        if _os.path.exists(name + '.bak'): _os.remove(name + '.bak')
        made.append(name)

    os.makedirs(MIRROR, exist_ok=True)
    for f in [pdf] + made:
        shutil.copy2(f, os.path.join(MIRROR, f))

    print(f'{len(pages)} pages + {len(apx)} appendix = {d.page_count} slides -> {pdf}')
    for m in made: print(f'  {m}')
    print(f'appendix labels : {", ".join(f"{k}={v}" for k, v in ref.items())}')
    print(f'font fallbacks  : {bad or "clean"}')
    print(f'text leaks      : {leaks or "clean"}')
    print(f'dangling refs   : {dangling or "clean"}')
    print(f'div balance     : {balance or "clean"}')
    if bad or leaks or dangling or balance:
        sys.exit('BUILD GATE FAILED')

if __name__ == '__main__':
    main()
