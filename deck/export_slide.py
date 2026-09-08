#!/usr/bin/env python3
"""Export one built slide on its own, for pasting into another deck.

    python3 export_slide.py 10            # by page number
    python3 export_slide.py 10 --name team

Writes three files into out_slides/, because Google Slides accepts a slide two
different ways and a PDF is useful for checking:
    <n>-<name>.pptx   one slide, 13.333x7.5in — use File > Import slides
    <n>-<name>.png    2x raster — use for a straight copy/paste
    <n>-<name>.pdf    single page

Run build.py first; this reads the PDF it produced.
"""
import os, sys, glob, re

DATE = '2026-09-09'
HERE = os.path.dirname(os.path.abspath(__file__))
OUT  = os.path.join(HERE, 'out_slides')


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    if not args:
        sys.exit(__doc__)
    n = int(args[0])
    name = ''
    if '--name' in sys.argv:
        name = sys.argv[sys.argv.index('--name') + 1]
    if not name:
        pages = sorted(glob.glob(os.path.join(HERE, 'pages/*.html')))
        if n <= len(pages):
            name = re.sub(r'^\d+-', '', os.path.splitext(os.path.basename(pages[n-1]))[0])
    stem = f'{n:02d}-{name}' if name else f'{n:02d}'

    import pymupdf
    pdf = os.path.join(HERE, f'Vishwa_Pitch_{DATE}.pdf')
    if not os.path.exists(pdf):
        sys.exit(f'{pdf} not found — run build.py first')
    d = pymupdf.open(pdf)
    if not 1 <= n <= d.page_count:
        sys.exit(f'page {n} is outside 1..{d.page_count}')

    os.makedirs(OUT, exist_ok=True)
    png = os.path.join(OUT, f'{stem}.png')
    d[n-1].get_pixmap(matrix=pymupdf.Matrix(2.0, 2.0)).save(png)

    one = pymupdf.open()
    one.insert_pdf(d, from_page=n-1, to_page=n-1)
    one.save(os.path.join(OUT, f'{stem}.pdf'))

    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    sl.shapes.add_picture(png, 0, 0, width=prs.slide_width, height=prs.slide_height)
    try:
        sys.path.insert(0, HERE)
        from notes import SLIDE_NOTES
        if SLIDE_NOTES.get(n):
            sl.notes_slide.notes_text_frame.text = SLIDE_NOTES[n].strip()
            print('speaker note attached')
    except Exception as e:
        print('note: no speaker note attached —', e)
    prs.save(os.path.join(OUT, f'{stem}.pptx'))

    print(f'page {n} -> out_slides/{stem}.{{pptx,png,pdf}}')


if __name__ == '__main__':
    main()
